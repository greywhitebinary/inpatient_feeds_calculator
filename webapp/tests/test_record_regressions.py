import os
import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from calculations import ons_delivery
from case_io import export_case_record_workbook, import_case_record_workbook
from data import (
    FORMULA_OPTIONAL_NUMERIC_COLUMNS,
    FORMULA_REQUIRED_COLUMNS,
    export_formulary_workbook,
    import_formulary_workbook,
    load_master_formulas,
    load_master_modulars,
    load_master_ons,
    validate_import,
)


class ServingOnsRecordRegressionTests(unittest.TestCase):
    def setUp(self):
        self.formulas = load_master_formulas().head(0)
        self.modulars = load_master_modulars().head(0)
        self.ons = (
            load_master_ons()
            .loc[lambda frame: frame["name"] == "BOOST Pudding — Vanilla"]
            .copy()
        )
        self.product_id = str(self.ons.iloc[0]["id"])
        self.servings_key = f"scenario_standard_ons_servings_{self.product_id}"
        self.times_key = f"scenario_standard_ons_times_{self.product_id}"

    def _payload_with_servings(self, value):
        state = {
            "scenario_standard_chosen_ons": ["BOOST Pudding — Vanilla"],
            self.servings_key: value,
            self.times_key: 2.0,
        }
        return export_case_record_workbook(
            state, self.formulas, self.modulars, self.ons
        )

    def test_boost_pudding_order_reopens_with_its_amounts_and_outcomes(self):
        payload = self._payload_with_servings(1.0)

        restored, _, _, restored_ons = import_case_record_workbook(BytesIO(payload))

        self.assertEqual(
            restored["scenario_standard_chosen_ons"],
            ["BOOST Pudding — Vanilla"],
        )
        self.assertEqual(restored[self.servings_key], 1.0)
        self.assertEqual(restored[self.times_key], 2.0)
        product = restored_ons.iloc[0].to_dict()
        self.assertEqual(product["calculation_basis"], "serving")
        self.assertEqual(product["serving_size_g"], 142)
        self.assertEqual(product["serving_unit"], "cup")

        outcomes = ons_delivery(
            product,
            containers_each_time=restored[self.servings_key],
            times_per_day=restored[self.times_key],
        )
        self.assertEqual(
            {
                key: outcomes[key]
                for key in (
                    "daily_servings",
                    "daily_volume_ml",
                    "energy_kcal",
                    "protein_g",
                    "carbohydrate_g",
                    "fat_g",
                    "fibre_g",
                    "free_water_ml",
                )
            },
            {
                "daily_servings": 2.0,
                "daily_volume_ml": 0.0,
                "energy_kcal": 460.0,
                "protein_g": 14.0,
                "carbohydrate_g": 64.0,
                "fat_g": 16.0,
                "fibre_g": 0.0,
                "free_water_ml": 186.0,
            },
        )

    def test_serving_quantity_rejects_invalid_types_and_ranges(self):
        invalid_values = ("one", True, -0.5, float("inf"))
        for value in invalid_values:
            with self.subTest(value=value):
                payload = self._payload_with_servings(value)
                with self.assertRaisesRegex(
                    ValueError,
                    r"(non-numeric|below 0|non-finite).*ons_servings",
                ):
                    import_case_record_workbook(BytesIO(payload))


class WorkbookLiteralTextRegressionTests(unittest.TestCase):
    def test_formula_like_case_label_is_literal_and_round_trips_exactly(self):
        label = '=HYPERLINK("https://example.invalid", "record")'
        with patch.dict(
            os.environ,
            {"CALCULATOR_WEBSITE_URL": "https://feeds.example.org/calculator"},
        ):
            payload = export_case_record_workbook(
                {"case_record_label": label},
                load_master_formulas().head(0),
                load_master_modulars().head(0),
            )

        workbook = load_workbook(BytesIO(payload), data_only=False)
        metadata = workbook["Case record"]
        self.assertEqual(metadata["B8"].value, label)
        self.assertEqual(metadata["B8"].data_type, "s")
        self.assertEqual(metadata["B6"].value, 1)
        self.assertEqual(metadata["B6"].data_type, "n")
        self.assertEqual(
            metadata["B2"].hyperlink.target,
            "https://feeds.example.org/calculator",
        )

        restored, _, _, _ = import_case_record_workbook(BytesIO(payload))
        self.assertEqual(restored["case_record_label"], label)

    def test_formula_like_product_text_is_literal_and_round_trips_exactly(self):
        formulas = load_master_formulas().iloc[[0]].copy()
        modulars = load_master_modulars().iloc[[0]].copy()
        ons = load_master_ons().iloc[[0]].copy()
        expected_names = (
            '=HYPERLINK("https://example.invalid/formula", "formula")',
            "=1+1",
            "=SUM(1, 2)",
        )
        formulas.loc[formulas.index[0], "name"] = expected_names[0]
        modulars.loc[modulars.index[0], "name"] = expected_names[1]
        ons.loc[ons.index[0], "name"] = expected_names[2]

        payload = export_formulary_workbook(formulas, modulars, ons)

        workbook = load_workbook(BytesIO(payload), data_only=False)
        for sheet_name, expected in zip(
            ("My Formulary", "My Modulars", "My ONS"), expected_names
        ):
            with self.subTest(sheet=sheet_name):
                sheet = workbook[sheet_name]
                columns = {
                    cell.value: cell.column
                    for cell in sheet[1]
                    if cell.value is not None
                }
                name_cell = sheet.cell(row=2, column=columns["name"])
                self.assertEqual(name_cell.value, expected)
                self.assertEqual(name_cell.data_type, "s")
        formula_sheet = workbook["My Formulary"]
        formula_columns = {
            cell.value: cell.column
            for cell in formula_sheet[1]
            if cell.value is not None
        }
        energy_cell = formula_sheet.cell(row=2, column=formula_columns["kcal_per_mL"])
        self.assertEqual(energy_cell.data_type, "n")

        restored_formulas, restored_modulars, restored_ons = import_formulary_workbook(
            BytesIO(payload)
        )
        self.assertEqual(restored_formulas.iloc[0]["name"], expected_names[0])
        self.assertEqual(restored_modulars.iloc[0]["name"], expected_names[1])
        self.assertEqual(restored_ons.iloc[0]["name"], expected_names[2])
        self.assertEqual(
            restored_formulas.iloc[0]["kcal_per_mL"],
            formulas.iloc[0]["kcal_per_mL"],
        )


class LiquidOnsSchemaRegressionTests(unittest.TestCase):
    def setUp(self):
        self.formulas = load_master_formulas().iloc[[0]].copy()
        self.modulars = load_master_modulars().iloc[[0]].copy()

    def test_blank_serving_numeric_fields_are_inapplicable_for_liquid_ons(self):
        ons = (
            load_master_ons()
            .loc[lambda frame: frame["calculation_basis"] == "container_ml"]
            .iloc[[0]]
            .copy()
        )
        serving_numeric_columns = {
            "serving_size_g",
            "kcal_per_serving",
            "protein_g_per_serving",
            "fat_g_per_serving",
            "carbohydrate_g_per_serving",
            "fibre_g_per_serving",
            "sodium_mg_per_serving",
            "potassium_mg_per_serving",
            "calcium_mg_per_serving",
            "magnesium_mg_per_serving",
            "phosphorus_mg_per_serving",
            "free_water_ml_per_serving",
        }
        ons[list(serving_numeric_columns)] = ons[list(serving_numeric_columns)].astype(
            object
        )
        ons.loc[:, list(serving_numeric_columns)] = None

        _, _, restored = validate_import(self.formulas, self.modulars, ons)

        for column in serving_numeric_columns:
            with self.subTest(column=column):
                self.assertEqual(restored.iloc[0][column], 0)

    def test_serving_ons_still_rejects_a_blank_required_nutrient(self):
        ons = (
            load_master_ons()
            .loc[lambda frame: frame["name"] == "BOOST Pudding — Vanilla"]
            .copy()
        )
        ons.loc[ons.index[0], "protein_g_per_serving"] = None

        with self.assertRaisesRegex(ValueError, "protein_g_per_serving"):
            validate_import(self.formulas, self.modulars, ons)

    def test_serving_ons_still_rejects_an_absent_required_nutrient(self):
        ons = (
            load_master_ons()
            .loc[lambda frame: frame["name"] == "BOOST Pudding — Vanilla"]
            .drop(columns=["protein_g_per_serving"])
            .copy()
        )

        with self.assertRaisesRegex(ValueError, "protein_g_per_serving"):
            validate_import(self.formulas, self.modulars, ons)


class OptionalFormulaColumnRegressionTests(unittest.TestCase):
    def setUp(self):
        self.formulas = load_master_formulas().iloc[[0]].copy()
        self.modulars = load_master_modulars().iloc[[0]].copy()
        self.ons = load_master_ons().iloc[[0]].copy()
        self.optional_only_columns = FORMULA_OPTIONAL_NUMERIC_COLUMNS - set(
            FORMULA_REQUIRED_COLUMNS
        )
        self.legacy_formulas = self.formulas.drop(
            columns=sorted(self.optional_only_columns)
        )

    def _assert_optional_columns_are_unknown(self, formulas):
        self.assertTrue(self.optional_only_columns.issubset(formulas.columns))
        for column in self.optional_only_columns:
            with self.subTest(column=column):
                self.assertTrue(pd.isna(formulas.iloc[0][column]))

    def test_legacy_formula_workbook_adds_missing_optional_columns_as_unknown(self):
        payload = export_formulary_workbook(
            self.legacy_formulas, self.modulars, self.ons
        )

        restored, _, _ = import_formulary_workbook(BytesIO(payload))

        self._assert_optional_columns_are_unknown(restored)
        for column in ("name", "kcal_per_mL", "protein_per_mL", "sodium_per_mL"):
            with self.subTest(preserved_column=column):
                self.assertEqual(
                    restored.iloc[0][column], self.legacy_formulas.iloc[0][column]
                )

    def test_added_unknown_columns_survive_a_second_workbook_round_trip(self):
        first_payload = export_formulary_workbook(
            self.legacy_formulas, self.modulars, self.ons
        )
        first_formulas, first_modulars, first_ons = import_formulary_workbook(
            BytesIO(first_payload)
        )

        second_payload = export_formulary_workbook(
            first_formulas, first_modulars, first_ons
        )
        second_formulas, _, _ = import_formulary_workbook(BytesIO(second_payload))

        self._assert_optional_columns_are_unknown(second_formulas)

    def test_master_loader_adds_missing_optional_columns_as_unknown(self):
        with patch("data.pd.read_csv", return_value=self.legacy_formulas):
            restored = load_master_formulas()

        self._assert_optional_columns_are_unknown(restored)

    def test_present_blank_fibre_keeps_the_declared_absence_convention(self):
        legacy = self.legacy_formulas.copy()
        legacy.loc[legacy.index[0], "fibre_per_mL"] = None

        restored, _, _ = validate_import(legacy, self.modulars, self.ons)

        self.assertEqual(restored.iloc[0]["fibre_per_mL"], 0)


if __name__ == "__main__":
    unittest.main()
